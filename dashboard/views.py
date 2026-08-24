import csv
import json
import threading
from decimal import Decimal
from datetime import timedelta

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from django import forms
from django.apps import apps
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Avg, Case, Count, DecimalField, ExpressionWrapper, F, Q, Sum, Value, When
from django.db.models.functions import Coalesce, TruncDate, TruncMonth
from django.forms import modelform_factory
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.html import strip_tags

from clientes.models import Cliente
from empresas.mail_utils import enviar_correo_saas, obtener_config_smtp
from empresas.mixins import requiere_modulo
from empresas.models import Empresa, SuscripcionEmpresa, Plan, Ticket, MensajeTicket, ConfiguracionSaaS
from productos.models import Producto
from tiendita_project.utils import form_errors_to_messages
from usuarios.models import AdminActionRequest, ModuloSaaS, RolEmpresa, UserProfile
from ventas.models import Venta, DetalleVenta

# Identificadores internos de roles (Mejor mantenibilidad)
ROLE_SUPERADMIN = 'superadmin'
ROLE_OWNER = 'owner'
ROLE_EMPLOYEE = 'employee'


DECIMAL_OUTPUT = DecimalField(max_digits=12, decimal_places=2)


def _valor_decimal(cantidad=0):
    return Value(Decimal(str(cantidad)), output_field=DECIMAL_OUTPUT)


def _construir_queryset_rentabilidad(empresa, fecha_inicio):
    """Base de detalles de venta para métricas de rentabilidad y rotación."""
    return (
        DetalleVenta.objects.filter(venta__empresa=empresa, venta__fecha__gte=fecha_inicio)
        .exclude(venta__estado='anulada')
        .annotate(fecha_venta=TruncDate('venta__fecha'))
        .annotate(
            utilidad_unitaria=Case(
                When(
                    producto__precio_costo__gt=0,
                    then=ExpressionWrapper(
                        F('precio_unitario') - F('producto__precio_costo'),
                        output_field=DECIMAL_OUTPUT,
                    ),
                ),
                default=Value(None, output_field=DECIMAL_OUTPUT),
                output_field=DECIMAL_OUTPUT,
            ),
            margen_unitario_pct=Case(
                When(
                    producto__precio_costo__gt=0,
                    then=ExpressionWrapper(
                        (F('precio_unitario') - F('producto__precio_costo')) * _valor_decimal(100) / F('producto__precio_costo'),
                        output_field=DECIMAL_OUTPUT,
                    ),
                ),
                default=Value(None, output_field=DECIMAL_OUTPUT),
                output_field=DECIMAL_OUTPUT,
            ),
        )
    )


def _calcular_metricas_rentabilidad(empresa, fecha_inicio, limite=10):
    """Agrupa métricas de utilidad, margen y rotación usando ORM."""
    detalles = _construir_queryset_rentabilidad(empresa, fecha_inicio)

    resumen = detalles.aggregate(
        rentabilidad_promedio=Coalesce(Avg('margen_unitario_pct'), _valor_decimal(0)),
        utilidad_total=Coalesce(Sum('utilidad_unitaria'), _valor_decimal(0)),
        unidades_vendidas=Count('id'),
        ventas_registradas=Count('venta_id', distinct=True),
    )

    top_rentabilidad_qs = (
        detalles.values('producto_id', 'producto__nombre', 'producto__tipo_producto__nombre')
        .annotate(
            margen_promedio=Avg('margen_unitario_pct'),
            utilidad_promedio=Avg('utilidad_unitaria'),
            utilidad_total=Sum('utilidad_unitaria'),
            unidades_vendidas=Count('id'),
        )
        .filter(margen_promedio__isnull=False)
        .order_by('-margen_promedio', '-utilidad_total', '-unidades_vendidas')[:limite]
    )

    rotacion_qs = (
        detalles.values('producto_id', 'producto__nombre', 'producto__tipo_producto__nombre')
        .annotate(
            ventas_count=Count('id'),
            dias_con_venta=Count('fecha_venta', distinct=True),
            utilidad_total=Sum('utilidad_unitaria'),
        )
        .annotate(
            rotacion_score=Case(
                When(
                    dias_con_venta__gt=0,
                    then=ExpressionWrapper(
                        F('ventas_count') * _valor_decimal(1) / F('dias_con_venta'),
                        output_field=DECIMAL_OUTPUT,
                    ),
                ),
                default=_valor_decimal(0),
                output_field=DECIMAL_OUTPUT,
            )
        )
        .filter(ventas_count__gt=0)
        .order_by('-rotacion_score', '-ventas_count', '-utilidad_total')
    )

    rotacion_media = rotacion_qs.aggregate(media=Avg('rotacion_score'))['media'] or Decimal('0')
    alta_rotacion_count = rotacion_qs.filter(rotacion_score__gte=rotacion_media).count()

    producto_mayor_rentabilidad = top_rentabilidad_qs.first()
    top_rotacion = list(rotacion_qs[:limite])

    categorias_rentables = (
        detalles.values('producto__categoria')
        .annotate(
            margen_promedio=Avg('margen_unitario_pct'),
            utilidad_total=Sum('utilidad_unitaria'),
            unidades_vendidas=Count('id'),
        )
        .filter(margen_promedio__isnull=False)
        .order_by('-utilidad_total', '-margen_promedio')[:limite]
    )

    return {
        'resumen': resumen,
        'producto_mayor_rentabilidad': producto_mayor_rentabilidad,
        'alta_rotacion_count': alta_rotacion_count,
        'top_rentabilidad': list(top_rentabilidad_qs),
        'top_rotacion': top_rotacion,
        'categorias_rentables': list(categorias_rentables),
    }

def home(request):
    """
    Redirige la ruta principal (/) al nuevo enrutador inteligente de SaaS.
    """
    return redirect('dashboard_router')


@login_required
def dashboard_router(request):
    """
    ESTRATEGIA DE DASHBOARDS: Enrutador Inteligente.
    Redirige al usuario a su ecosistema correspondiente según su rol.
    """
    usuario = request.user
    perfil = getattr(usuario, 'profile', None)
    
    # El onboarding es estrictamente para los DUEÑOS
    if perfil and perfil.rol == ROLE_OWNER and not perfil.onboarding_completado:
        return redirect('onboarding')
        
    if perfil and perfil.rol == ROLE_SUPERADMIN:
        return redirect('dashboard_saas')
    elif perfil and perfil.rol in [ROLE_OWNER, ROLE_EMPLOYEE]:
        return redirect('dashboard_tenant')
    else:
        # Si es cliente final, va al catálogo de su empresa
        if perfil and getattr(perfil, 'empresa', None):
            return redirect('tienda_publica', slug=perfil.empresa.slug)
        return redirect('home')  # Fallback


@login_required
def dashboard_saas(request):
    """Dashboard GLOBAL y exclusivo para el dueño del software (SuperAdmin)"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales del sistema.")
        
    # Métricas Globales SaaS
    total_empresas = Empresa.objects.count()
    empresas_activas = Empresa.objects.filter(activa=True).count()
    
    suscripciones = SuscripcionEmpresa.objects.select_related('empresa', 'plan').all()
    mrr_estimado = suscripciones.aggregate(total=Sum('plan__precio_mensual'))['total'] or 0
    
    # --- DATOS REALES PARA GRÁFICAS SaaS ---
    # Distribución real de empresas activas por plan
    planes_stats = Plan.objects.annotate(
        num_suscripciones=Count('suscripciones', filter=Q(suscripciones__empresa__activa=True))
    ).values('nombre', 'num_suscripciones')
    
    planes_labels = [p['nombre'] for p in planes_stats] if planes_stats else ['Sin Planes']
    planes_series = [p['num_suscripciones'] for p in planes_stats] if planes_stats else [1]
    # ----------------------------------------

    context = {
        'total_empresas': total_empresas,
        'empresas_activas': empresas_activas,
        'suscripciones': suscripciones,
        'mrr_estimado': mrr_estimado,
        'planes_labels': json.dumps(planes_labels),
        'planes_series': json.dumps(planes_series),
    }
    return render(request, 'dashboard/saas_dashboard.html', context)


@login_required
def saas_empresas(request):
    """Vista del SuperAdmin para gestionar Inquilinos (Tenants)"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    empresas = Empresa.objects.select_related('suscripcion').prefetch_related('empleados__user').order_by('-fecha_creacion')
    return render(request, 'dashboard/saas_empresas.html', {'empresas': empresas})


@login_required
def exportar_empresas_csv(request):
    """Exporta el listado de empresas a CSV para el Super Admin"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    filename = f'empresas_saas_{timezone.now().strftime("%Y%m%d")}.csv'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    response.write(u'\ufeff'.encode('utf8'))  # Forzar codificación UTF-8 para Excel
    writer = csv.writer(response)
    writer.writerow(['ID', 'Nombre de Empresa', 'URL Slug', 'Estado', 'Fecha Registro', 'Teléfono', 'Email de Contacto'])
    
    # Prefetch empleados para sacar el email del dueño de forma eficiente
    empresas = Empresa.objects.prefetch_related('empleados__user').all()
    
    for emp in empresas:
        estado = "Activa" if emp.activa else "Suspendida"
        dueño_perfil = next((p for p in emp.empleados.all() if p.rol == ROLE_OWNER), None)
        email = dueño_perfil.user.email if dueño_perfil and dueño_perfil.user else "Sin email"
        
        writer.writerow([emp.id, emp.nombre, emp.slug, estado, emp.fecha_creacion.strftime("%Y-%m-%d"), emp.telefono_whatsapp, email])
        
    return response


@login_required
def toggle_estado_empresa(request, empresa_id):
    """Suspende o activa una empresa (Tenant)"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN or request.method != 'POST':
        raise PermissionDenied("Acción no permitida.")
        
    empresa = get_object_or_404(Empresa, id=empresa_id)
    empresa.activa = not empresa.activa
    empresa.save()
    
    if empresa.activa:
        # Si se está reactivando, verificar si la suscripción está vencida
        if hasattr(empresa, 'suscripcion') and empresa.suscripcion.dias_restantes() <= 0:
            # Otorgar 24 horas de gracia para que puedan entrar a pagar o usar el sistema
            empresa.suscripcion.fecha_vencimiento = timezone.now() + timedelta(days=1)
            empresa.suscripcion.save()
            messages.warning(request, f"Empresa '{empresa.nombre}' activada. Su suscripción estaba vencida, se le otorgó 1 día de gracia.")
        else:
            messages.success(request, f"La empresa '{empresa.nombre}' ha sido activada con éxito.")
    else:
        messages.success(request, f"La empresa '{empresa.nombre}' ha sido suspendida con éxito.")
        
    return redirect('saas_empresas')


@login_required
def saas_configuracion(request):
    """Vista del SuperAdmin para configuración global del SaaS"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    # Obtener la configuración o crearla si es la primera vez (singleton)
    config, created = ConfiguracionSaaS.objects.get_or_create(id=1)
        
    if request.method == 'POST':
        # Branding y Contacto
        config.nombre_plataforma = request.POST.get('app_name', config.nombre_plataforma)
        config.support_email = request.POST.get('support_email', config.support_email)
        config.telefono_whatsapp = request.POST.get('telefono_whatsapp', config.telefono_whatsapp)
        config.instagram_url = request.POST.get('instagram_url', config.instagram_url)
        config.facebook_url = request.POST.get('facebook_url', config.facebook_url)
        config.sitio_web = request.POST.get('sitio_web', config.sitio_web)
        config.tiktok_url = request.POST.get('tiktok_url', config.tiktok_url)
        config.x_url = request.POST.get('x_url', config.x_url)

        # Configuración SMTP
        config.smtp_host = request.POST.get('smtp_host', '')
        config.smtp_port = request.POST.get('smtp_port', 587) or 587
        config.smtp_user = request.POST.get('smtp_user', '')
        smtp_pass = request.POST.get('smtp_pass', '')
        if smtp_pass:
            config.smtp_password = smtp_pass
        config.smtp_encryption = request.POST.get('smtp_encryption', 'starttls')
        config.smtp_from_name = request.POST.get('smtp_from_name', '')
        config.smtp_from_email = request.POST.get('smtp_from_email', '')
        config.smtp_test_mode = request.POST.get('smtp_test_mode') == 'on'
        config.smtp_test_recipient = request.POST.get('smtp_test_recipient', '')
        config.smtp_max_attachment_mb = request.POST.get('smtp_max_attachment_mb', 25) or 25
        config.smtp_daily_limit = request.POST.get('smtp_daily_limit', 500) or 500
        config.smtp_retry_attempts = request.POST.get('smtp_retry_attempts', 3) or 3
        config.smtp_retry_delay_seconds = request.POST.get('smtp_retry_delay_seconds', 2) or 2
        config.smtp_last_test_status = 'pending'
        config.smtp_last_test_message = 'Configuración guardada. Pendiente de verificación.'
        config.smtp_last_test_at = timezone.now()
        
        # Pasarelas de Pago
        stripe_secret = request.POST.get('stripe_secret', '')
        if stripe_secret:
            config.stripe_secret_key = stripe_secret
            
        config.paypal_client_id = request.POST.get('paypal_client_id', '')
        paypal_secret = request.POST.get('paypal_secret', '')
        if paypal_secret:
            config.paypal_secret_key = paypal_secret
            
        payphone_token = request.POST.get('payphone_token', '')
        if payphone_token:
            config.payphone_token = payphone_token

        config.mercadopago_active = request.POST.get('mercadopago_active') == 'on'
        mercadopago_token = request.POST.get('mercadopago_token', '')
        if mercadopago_token:
            config.mercadopago_token = mercadopago_token

        config.coinbase_active = request.POST.get('coinbase_active') == 'on'
        coinbase_api_key = request.POST.get('coinbase_api_key', '')
        if coinbase_api_key:
            config.coinbase_api_key = coinbase_api_key

        config.razorpay_active = request.POST.get('razorpay_active') == 'on'
        config.razorpay_key_id = request.POST.get('razorpay_key_id', '')
        razorpay_secret = request.POST.get('razorpay_key_secret', '')
        if razorpay_secret:
            config.razorpay_key_secret = razorpay_secret

        # Transferencia Bancaria
        config.transferencia_active = request.POST.get('transferencia_active') == 'on'
        config.banco_pichincha = request.POST.get('banco_pichincha', '')
        config.banco_pichincha_2 = request.POST.get('banco_pichincha_2', '')
        config.banco_guayaquil = request.POST.get('banco_guayaquil', '')
        config.banco_jep = request.POST.get('banco_jep', '')

        # NOTA: Los siguientes campos no existen en el modelo ConfiguracionSaaS y causarían un error.
        # config.theme_color, config.stripe_active, config.stripe_public_key, config.paypal_active,
        # config.payphone_active, config.force_2fa, config.session_timeout

        config.save()
        
        cache.delete('saas_app_name')

        tab_activa = request.POST.get('active_tab', 'branding')
        messages.success(request, 'Ajustes del sistema guardados correctamente.')
        return redirect(f"{reverse('saas_configuracion')}#{tab_activa}")
        
    return render(request, 'dashboard/saas_configuracion.html', {'config': config})


@login_required
def saas_configuracion_test_smtp(request):
    """Prueba el SMTP actual o el SMTP capturado desde el formulario sin guardar."""
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'message': 'Método no permitido.'}, status=405)

    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        return JsonResponse({'ok': False, 'message': 'Acceso exclusivo para administradores globales.'}, status=403)

    config, _ = ConfiguracionSaaS.objects.get_or_create(id=1)
    overrides = {
        'smtp_host': request.POST.get('smtp_host', '').strip(),
        'smtp_port': request.POST.get('smtp_port', '').strip(),
        'smtp_user': request.POST.get('smtp_user', '').strip(),
        'smtp_password': request.POST.get('smtp_pass', '').strip(),
        'smtp_encryption': request.POST.get('smtp_encryption', 'starttls').strip(),
        'smtp_from_name': request.POST.get('smtp_from_name', '').strip(),
        'smtp_from_email': request.POST.get('smtp_from_email', '').strip(),
        'smtp_test_mode': request.POST.get('smtp_test_mode') == 'on',
        'smtp_test_recipient': request.POST.get('smtp_test_recipient', '').strip(),
        'smtp_retry_attempts': request.POST.get('smtp_retry_attempts', '').strip(),
        'smtp_retry_delay_seconds': request.POST.get('smtp_retry_delay_seconds', '').strip(),
        'smtp_max_attachment_mb': request.POST.get('smtp_max_attachment_mb', '').strip(),
        'smtp_daily_limit': request.POST.get('smtp_daily_limit', '').strip(),
    }

    recipient = overrides['smtp_test_recipient'] or request.user.email or config.support_email
    if not recipient:
        return JsonResponse({'ok': False, 'message': 'Debes indicar un correo de prueba o tener un correo en tu usuario.'}, status=400)

    subject = f"Prueba SMTP - {config.nombre_plataforma or 'Tiendita SaaS'}"
    body = (
        f"Esta es una prueba de correo desde {config.nombre_plataforma or 'Tiendita SaaS'}.\n\n"
        f"Servidor: {overrides['smtp_host'] or config.smtp_host or '(sin definir)'}\n"
        f"Puerto: {overrides['smtp_port'] or config.smtp_port or '(sin definir)'}\n"
        f"Cifrado: {overrides['smtp_encryption'] or config.smtp_encryption or 'starttls'}\n"
        f"Remitente: {overrides['smtp_from_email'] or config.smtp_from_email or '(sin definir)'}\n"
    )

    try:
        smtp_config_used = obtener_config_smtp(config=config, overrides=overrides)

        html_message = f"""
        <html>
            <body style="font-family: Arial, sans-serif; background:#f8fafc; padding:24px; color:#0f172a;">
                <div style="max-width:560px; margin:0 auto; background:#fff; border:1px solid #e2e8f0; border-radius:16px; padding:28px;">
                    <h2 style="margin:0 0 12px; color:#2563eb;">{config.nombre_plataforma or 'Tiendita SaaS'}</h2>
                    <h3 style="margin:0 0 16px; color:#111827;">Prueba SMTP exitosa</h3>
                    <p style="margin:0 0 12px; color:#475569;">Este correo confirma que la configuración SMTP respondió correctamente.</p>
                    <ul style="padding-left:20px; color:#334155;">
                        <li>Servidor: {overrides['smtp_host'] or config.smtp_host or '(sin definir)'}</li>
                        <li>Puerto: {overrides['smtp_port'] or config.smtp_port or '(sin definir)'}</li>
                        <li>Cifrado: {overrides['smtp_encryption'] or config.smtp_encryption or 'starttls'}</li>
                        <li>Remitente: {overrides['smtp_from_email'] or config.smtp_from_email or config.support_email or '(sin definir)'}</li>
                        <li>Destino de prueba: {recipient}</li>
                    </ul>
                </div>
            </body>
        </html>
        """
        enviados = enviar_correo_saas(
            subject,
            body,
            [recipient],
            html_message=html_message,
            config=config,
            overrides=overrides,
            fail_silently=False,
        )
        config.smtp_last_test_status = 'success'
        success_message = f'¡Conexión exitosa! Se envió un correo de prueba a {recipient}.'
        config.smtp_last_test_message = success_message
        config.smtp_last_test_at = timezone.now()
        config.save(update_fields=['smtp_last_test_status', 'smtp_last_test_message', 'smtp_last_test_at'])

        debug_info = {
            "host": smtp_config_used['host'],
            "port": smtp_config_used['port'],
            "user": smtp_config_used['username'],
            "from": smtp_config_used['from_email'],
            "to": recipient,
            "encryption": smtp_config_used['encryption'],
        }

        return JsonResponse({'ok': True, 'message': success_message, 'debug_info': debug_info})
    except Exception as exc:
        config.smtp_last_test_status = 'error'
        config.smtp_last_test_message = str(exc)
        config.smtp_last_test_at = timezone.now()
        config.save(update_fields=['smtp_last_test_status', 'smtp_last_test_message', 'smtp_last_test_at'])
        return JsonResponse({'ok': False, 'message': f'No se pudo enviar el correo de prueba: {exc}'}, status=400)


@login_required
def saas_tickets(request):
    """Bandeja de entrada (Inbox) de Soporte Técnico para el SuperAdmin"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    ticket_id = request.GET.get('id')
    filtro_actual = request.GET.get('filtro', 'inbox')
    ticket_activo = None
    
    if request.method == 'POST' and ticket_id:
        ticket = get_object_or_404(Ticket, id=ticket_id)
        
        if 'eliminar_ticket' in request.POST and request.POST.get('eliminar_ticket') == '1':
            ticket.delete()
            messages.success(request, "El ticket fue eliminado de forma permanente.")
            return redirect(f"{reverse('saas_tickets')}?filtro={filtro_actual}")
            
        mensaje_texto = request.POST.get('mensaje', '').strip()
        nuevo_estado = request.POST.get('estado')
        adjunto = request.FILES.get('adjunto')
        
        enviar_correo = False
        
        if mensaje_texto or adjunto:
            MensajeTicket.objects.create(ticket=ticket, remitente=request.user, mensaje=mensaje_texto, adjunto=adjunto, es_staff=True)
            enviar_correo = True
        
        if nuevo_estado and nuevo_estado != ticket.estado:
            ticket.estado = nuevo_estado
            ticket.save()
            if nuevo_estado == 'resuelto':
                messages.success(request, "Ticket archivado correctamente.")
            else:
                messages.success(request, "Estado del ticket actualizado.")
            
        # ENVÍO DE CORREO AUTOMÁTICO - Respuesta del Equipo de Soporte
        if enviar_correo and ticket.usuario and ticket.usuario.email:
            def enviar_respuesta_async():
                try:
                    mensaje_txt = f"Hola {ticket.usuario.username},\n\nEl equipo de soporte ha respondido a tu ticket #{ticket.id} ({ticket.asunto}):\n\n{mensaje_texto}\n\nSaludos,\nEquipo de Soporte"
                    
                    app_name = cache.get('saas_app_name', 'Tiendita SaaS')
                    mensaje_html = render_to_string('emails/ticket_respuesta.html', {
                        'app_name': app_name, 'ticket_id': ticket.id, 'usuario': ticket.usuario.username, 'asunto': ticket.asunto, 'respuesta': mensaje_texto
                    })
                    
                    enviar_correo_saas(
                        f"Respuesta a tu ticket #{ticket.id}",
                        mensaje_txt,
                        [ticket.usuario.email],
                        html_message=mensaje_html,
                        fail_silently=True,
                    )
                except Exception as e:
                    print(f"Error enviando respuesta de ticket: {e}")
            threading.Thread(target=enviar_respuesta_async).start()
                
        return redirect(f"{reverse('saas_tickets')}?id={ticket.id}&filtro={filtro_actual}")
        
    if filtro_actual == 'archivados':
        tickets = Ticket.objects.filter(estado='resuelto').select_related('empresa').prefetch_related('mensajes').order_by('-actualizado_en')
    else:
        tickets = Ticket.objects.exclude(estado='resuelto').select_related('empresa').prefetch_related('mensajes').order_by('-actualizado_en')
        
    if ticket_id:
        ticket_activo = get_object_or_404(Ticket, id=ticket_id)
        
    return render(request, 'dashboard/saas_tickets.html', {'tickets': tickets, 'ticket_activo': ticket_activo, 'filtro_actual': filtro_actual})


@login_required
def saas_control_panel(request):
    """Hub central de la base de datos para crear/editar todos los modelos (CRUD)"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    # Estadísticas para el header del panel
    context = {
        'total_usuarios': User.objects.count(),
        'total_planes': Plan.objects.count(),
        'total_modulos': ModuloSaaS.objects.count(),
        'total_roles': RolEmpresa.objects.count(),
        'total_empresas': Empresa.objects.count(),
    }
    
    return render(request, 'dashboard/saas_database.html', context)


@login_required
def saas_crear_registro(request, app_label, model_name):
    """Fábrica de Formularios: Crea dinámicamente el formulario para CUALQUIER modelo"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")

    try:
        Modelo = apps.get_model(app_label, model_name)
    except LookupError:
        messages.error(request, "El módulo o tabla solicitada no existe.")
        return redirect('saas_control_panel')

    es_modelo_user = Modelo is User

    if es_modelo_user:
        # CASO ESPECIAL: el modelo User de Django necesita tratamiento distinto.
        # - 'password' no debe guardarse en texto plano: hay que usar set_password().
        # - 'groups', 'user_permissions', 'last_login', 'date_joined' no aplican
        #   a este SaaS (los permisos reales se gestionan vía UserProfile.rol).
        campos_user = ['username', 'email', 'first_name', 'last_name', 'is_active', 'is_staff', 'is_superuser']
        FormularioDinamico = modelform_factory(Modelo, fields=campos_user)
    else:
        FormularioDinamico = modelform_factory(Modelo, exclude=[])

    if request.method == 'POST':
        form = FormularioDinamico(request.POST, request.FILES)
        password_nueva = request.POST.get('password', '').strip() if es_modelo_user else None

        if es_modelo_user and not password_nueva:
            messages.error(request, "Debes asignar una contraseña para el nuevo usuario.")
            form.add_error(None, "La contraseña es obligatoria.")

        if form.is_valid() and not (es_modelo_user and not password_nueva):
            instancia = form.save(commit=False)
            if es_modelo_user:
                instancia.set_password(password_nueva)
            instancia.save()
            if hasattr(form, 'save_m2m'):
                form.save_m2m()
            messages.success(request, f"¡Registro creado exitosamente en {Modelo._meta.verbose_name.title()}!")
            return redirect('saas_control_panel')
        else:
            form_errors_to_messages(request, form)
    else:
        form = FormularioDinamico()

    # Inyectar estilos Premium (Bootstrap) a todos los campos dinámicamente
    for field_name, field in form.fields.items():
        if isinstance(field.widget, forms.CheckboxInput):
            field.widget.attrs.update({'class': 'form-check-input ms-0 me-2', 'style': 'width: 1.5em; height: 1.5em; cursor: pointer;'})
        elif isinstance(field.widget, forms.SelectMultiple):
            # Un <select multiple> nativo requiere Ctrl/Cmd+Click para elegir
            # varios sin perder la selección anterior — muy poco intuitivo.
            # Checkboxes: cada opción se marca/desmarca con un solo click.
            nuevo_widget = forms.CheckboxSelectMultiple(attrs={'class': 'form-check-input'})
            nuevo_widget.choices = field.widget.choices  # heredar las opciones del widget original
            field.widget = nuevo_widget
        elif isinstance(field.widget, forms.Select):
            field.widget.attrs.update({'class': 'form-select form-select-lg bg-light border-0 shadow-sm', 'style': 'border-radius: 12px;'})
        elif isinstance(field.widget, forms.Textarea):
            field.widget.attrs.update({'class': 'form-control bg-light border-0 shadow-sm', 'style': 'border-radius: 12px;', 'rows': 3})
        else:
            field.widget.attrs.update({'class': 'form-control form-control-lg bg-light border-0 shadow-sm', 'style': 'border-radius: 12px;'})

    context = {
        'form': form,
        'titulo': f"Añadir {Modelo._meta.verbose_name.title()}",
        'es_modelo_user': es_modelo_user,
        'es_creacion': True,
    }
    return render(request, 'dashboard/saas_generic_form.html', context)


@login_required
def saas_listar_registros(request, app_label, model_name):
    """Fábrica de Vistas: Lista dinámicamente CUALQUIER modelo del sistema"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")

    try:
        Modelo = apps.get_model(app_label, model_name)
    except LookupError:
        messages.error(request, "El módulo o tabla solicitada no existe.")
        return redirect('saas_control_panel')

    query = request.GET.get('q', '').strip()
    
    # Buscador Mágico: Busca en todos los campos de texto
    if query:
        q_objects = Q()
        for field in Modelo._meta.fields:
            if field.get_internal_type() in ['CharField', 'TextField']:
                q_objects |= Q(**{f"{field.name}__icontains": query})
        
        registros = Modelo.objects.filter(q_objects).order_by('-id')[:100] if q_objects else Modelo.objects.all().order_by('-id')[:100]
    else:
        registros = Modelo.objects.all().order_by('-id')[:100]
    
    # Obtener los campos para la tabla (excluyendo contraseñas o campos conflictivos M2M)
    campos = [f for f in Modelo._meta.fields if f.name not in ['password']]
        
    campos = campos[:6] # Limitar a un máximo de 6 columnas para no desbordar la pantalla
    
    registros_data = []
    for reg in registros:
        fila = []
        for campo in campos:
            # Obtener el valor legible (ej: en lugar de 'en_progreso', mostrar 'En Progreso')
            if hasattr(reg, f'get_{campo.name}_display'):
                valor = getattr(reg, f'get_{campo.name}_display')()
            else:
                valor = getattr(reg, campo.name)
            fila.append(valor)
        registros_data.append({'id': reg.id, 'fila': fila})

    context = {
        'titulo': f"Directorio de {Modelo._meta.verbose_name_plural.title()}",
        'modelo_nombre': Modelo._meta.verbose_name.title(),
        'app_label': app_label,
        'model_name': model_name,
        'campos': campos,
        'registros_data': registros_data,
    }
    return render(request, 'dashboard/saas_generic_list.html', context)


@login_required
def saas_editar_registro(request, app_label, model_name, object_id):
    """Fábrica de Edición: Edita cualquier registro existente de manera automática"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")

    try:
        Modelo = apps.get_model(app_label, model_name)
    except LookupError:
        messages.error(request, "El módulo o tabla solicitada no existe.")
        return redirect('saas_control_panel')

    instancia = get_object_or_404(Modelo, id=object_id)
    es_modelo_user = Modelo is User

    if es_modelo_user:
        campos_user = ['username', 'email', 'first_name', 'last_name', 'is_active', 'is_staff', 'is_superuser']
        FormularioDinamico = modelform_factory(Modelo, fields=campos_user)
    else:
        FormularioDinamico = modelform_factory(Modelo, exclude=[])

    if request.method == 'POST':
        form = FormularioDinamico(request.POST, request.FILES, instance=instancia)
        password_nueva = request.POST.get('password', '').strip() if es_modelo_user else None

        if form.is_valid():
            instancia_guardada = form.save(commit=False)
            if es_modelo_user and password_nueva:
                # Solo se cambia la contraseña si el admin escribió algo en el campo.
                instancia_guardada.set_password(password_nueva)
            instancia_guardada.save()
            if hasattr(form, 'save_m2m'):
                form.save_m2m()
            messages.success(request, f"¡Registro actualizado exitosamente en {Modelo._meta.verbose_name.title()}!")
            return redirect('saas_listar_registros', app_label=app_label, model_name=model_name)
        else:
            form_errors_to_messages(request, form)
    else:
        form = FormularioDinamico(instance=instancia)

    # Aplicar estilos Bootstrap
    for field_name, field in form.fields.items():
        if isinstance(field.widget, forms.CheckboxInput):
            field.widget.attrs.update({'class': 'form-check-input ms-0 me-2', 'style': 'width: 1.5em; height: 1.5em; cursor: pointer;'})
        elif isinstance(field.widget, forms.SelectMultiple):
            nuevo_widget = forms.CheckboxSelectMultiple(attrs={'class': 'form-check-input'})
            nuevo_widget.choices = field.widget.choices  # heredar las opciones del widget original
            field.widget = nuevo_widget
        elif isinstance(field.widget, forms.Select) or isinstance(field.widget, forms.Textarea) or isinstance(field.widget, forms.TextInput):
            field.widget.attrs.update({'class': 'form-control form-control-lg bg-light border-0 shadow-sm', 'style': 'border-radius: 12px;'})

    context = {
        'form': form,
        'titulo': f"Editar {Modelo._meta.verbose_name.title()} #{object_id}",
        'es_modelo_user': es_modelo_user,
        'es_creacion': False,
    }
    return render(request, 'dashboard/saas_generic_form.html', context)


@login_required
def saas_eliminar_registro(request, app_label, model_name, object_id):
    """Fábrica de Eliminación: Borra dinámicamente un registro"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")

    try:
        Modelo = apps.get_model(app_label, model_name)
    except LookupError:
        messages.error(request, "El módulo o tabla solicitada no existe.")
        return redirect('saas_control_panel')

    instancia = get_object_or_404(Modelo, id=object_id)
    if request.method == 'POST':
        instancia.delete()
        messages.success(request, f"Registro eliminado exitosamente.")
        
    return redirect('saas_listar_registros', app_label=app_label, model_name=model_name)


@login_required
def dashboard_tenant(request):
    """Dashboard OPERATIVO exclusivo para la Empresa (Dueño y Empleados)"""
    empresa = getattr(request, 'empresa', None)
    if not empresa:
        return redirect('onboarding')

    es_dueno = getattr(getattr(request.user, 'profile', None), 'rol', '') == ROLE_OWNER
        
    hoy = timezone.now()
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0)
    
    ventas_mes = Venta.objects.filter(
        empresa=empresa, fecha__gte=inicio_mes
    ).exclude(estado='anulada').aggregate(Sum('total'))['total__sum'] or 0
    productos_activos = Producto.objects.filter(empresa=empresa, estado='disponible').count()
    nuevos_clientes = Cliente.objects.filter(empresas=empresa, creado_en__gte=inicio_mes).count()
    
    hace_30_dias = hoy - timedelta(days=29) # 30 días incluyendo hoy
    ventas_diarias = Venta.objects.filter(
        empresa=empresa, 
        fecha__gte=hace_30_dias
    ).exclude(estado='anulada').values('fecha__date').annotate(total_dia=Sum('total')).order_by('fecha__date')
    
    # Convertir a diccionario para llenado rápido
    ventas_dict = {v['fecha__date']: float(v['total_dia']) for v in ventas_diarias}
    
    fechas_chart = []
    totales_chart = []
    
    # Iterar los últimos 30 días garantizando que no haya huecos en la gráfica
    for i in range(29, -1, -1):
        fecha_iter = (hoy - timedelta(days=i)).date()
        fechas_chart.append(fecha_iter.strftime('%d %b'))
        totales_chart.append(ventas_dict.get(fecha_iter, 0.0))
    # -------------------------------------------------------------

    metricas_rentabilidad = _calcular_metricas_rentabilidad(empresa, hoy - timedelta(days=29)) if es_dueno else {}

    context = {
        'empresa': empresa,
        'ventas_mes': ventas_mes,
        'productos_activos': productos_activos,
        'nuevos_clientes': nuevos_clientes,
        'fechas_chart': json.dumps(fechas_chart),
        'totales_chart': json.dumps(totales_chart),
        'es_dueno': es_dueno,
        'metricas_rentabilidad': metricas_rentabilidad,
    }
    return render(request, 'dashboard/tenant_dashboard.html', context)


@login_required
@requiere_modulo('reportes')
def tenant_reportes(request):
    """Módulo Avanzado de Reportes y Estadísticas para el Dueño"""
    empresa = getattr(request, 'empresa', None)
    if not empresa:
        return redirect('onboarding')

    es_dueno = getattr(getattr(request.user, 'profile', None), 'rol', '') == ROLE_OWNER
        
    # Filtro de días (Predeterminado a 30)
    try:
        dias = int(request.GET.get('dias', 30))
    except ValueError:
        dias = 30
        
    fecha_inicio = timezone.now() - timedelta(days=dias)
    
    ventas_periodo = Venta.objects.filter(empresa=empresa, fecha__gte=fecha_inicio).select_related('cliente')
    metricas_rentabilidad = _calcular_metricas_rentabilidad(empresa, fecha_inicio) if es_dueno else {}
    
    tipo_exportacion = request.GET.get('export')
    if tipo_exportacion in ['csv', 'excel']:
        if tipo_exportacion == 'excel':
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="reporte_ventas_{dias}_dias.xlsx"'
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Ventas"
            
            # --- TÍTULOS SUPERIORES ---
            ws.append([f"REPORTE DE VENTAS - {empresa.nombre.upper()}"])
            ws.append([f"Período: Últimos {dias} días"])
            ws.append([])
            
            # --- CABECERAS DE TABLA ---
            headers = ['ID Venta', 'Fecha', 'Cliente', 'Subtotal', 'Descuento', 'Total', 'Estado', 'Origen']
            ws.append(headers)
            
            # Estilos de Títulos y Cabeceras
            ws['A1'].font = Font(size=14, bold=True, color="1F2937")
            ws['A2'].font = Font(italic=True, color="6B7280")
            
            fill_header = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            font_header = Font(color="FFFFFF", bold=True)
            thin_border = Border(left=Side(style='thin', color='4F46E5'), right=Side(style='thin', color='4F46E5'), top=Side(style='thin', color='4F46E5'), bottom=Side(style='thin', color='4F46E5'))
            
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=4, column=col_num)
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border
                
            # Estilos estrictos para las celdas de datos
            fill_blanco = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            fill_gris = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
            borde_datos = Border(
                left=Side(style='thin', color='E5E7EB'), right=Side(style='thin', color='E5E7EB'), 
                top=Side(style='thin', color='E5E7EB'), bottom=Side(style='thin', color='E5E7EB')
            )
            align_datos = Alignment(horizontal='left', vertical='center', wrap_text=False)
            align_centro = Alignment(horizontal='center', vertical='center', wrap_text=False)
            
            for row_idx, v in enumerate(ventas_periodo.order_by('-fecha'), start=5):
                cliente_nombre = v.cliente.nombre if v.cliente else 'Consumidor Final'
                
                try:
                    estado = v.get_estado_display()
                except AttributeError:
                    estado = getattr(v, 'estado', 'Completada')
                try:
                    origen = v.get_origen_display()
                except AttributeError:
                    origen = getattr(v, 'origen', 'Manual')
                    
                data_row = [
                    v.id, v.fecha.strftime("%Y-%m-%d %H:%M"), cliente_nombre, 
                    float(v.subtotal) if v.subtotal else 0, 
                    float(v.descuento) if v.descuento else 0, 
                    float(v.total) if v.total else 0, 
                    estado, origen
                ]
                
                # Inyectamos una barrera invisible (espacio) al final de la fila
                # para forzar a Excel a recortar la última columna y evitar desbordamientos
                data_row.append(" ")
                ws.append(data_row)
                
                # Aplicar color alternado (Zebra Striping)
                current_fill = fill_blanco if row_idx % 2 == 1 else fill_gris
                
                # Aplicar estilos, bordes y alineación a TODAS las celdas de la fila
                for col_num in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_num)
                    cell.fill = current_fill
                    cell.border = borde_datos
                    if col_num in [4, 5, 6]:
                        cell.number_format = '"$"#,##0.00'
                        cell.alignment = align_centro
                    elif col_num in [1, 2, 7, 8]:
                        cell.alignment = align_centro
                    else:
                        cell.alignment = align_datos
                        
            # Aplicar anchos estrictos y controlados para la tabla de Ventas
            anchos = {'A': 10, 'B': 18, 'C': 22, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 16}
            for col, ancho in anchos.items():
                ws.column_dimensions[col].width = ancho
                
            wb.save(response)
            return response
        else:
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = f'attachment; filename="reporte_ventas_{dias}_dias.csv"'
            response.write(u'\ufeff'.encode('utf8'))  # BOM para tildes
            writer = csv.writer(response)
            
            # --- TÍTULOS SUPERIORES CSV ---
            writer.writerow([f"REPORTE DE VENTAS - {empresa.nombre.upper()}"])
            writer.writerow([f"Período: Últimos {dias} días"])
            writer.writerow([])
            writer.writerow(['ID Venta', 'Fecha', 'Cliente', 'Subtotal', 'Descuento', 'Total', 'Estado', 'Origen'])
        
        for v in ventas_periodo.order_by('-fecha'):
            cliente_nombre = v.cliente.nombre if v.cliente else 'Consumidor Final'
            
            try:
                estado = v.get_estado_display()
            except AttributeError:
                estado = getattr(v, 'estado', 'Completada')
            try:
                origen = v.get_origen_display()
            except AttributeError:
                origen = getattr(v, 'origen', 'Manual')
                
            writer.writerow([
                v.id, 
                v.fecha.strftime("%Y-%m-%d %H:%M"), 
                cliente_nombre, 
                v.subtotal, 
                v.descuento, 
                v.total, 
                estado,
                origen
            ])
        return response
        
    # Filtrar ventas anuladas para los cálculos financieros
    ventas_validas = ventas_periodo.exclude(estado='anulada')

    # 2. KPIs Rápidos
    total_ingresos = ventas_validas.aggregate(Sum('total'))['total__sum'] or 0
    total_ventas = ventas_validas.count()
    ticket_promedio = (total_ingresos / total_ventas) if total_ventas > 0 else 0
    
    # 3. Gráfica de Ingresos Organizados y Rellenos
    fechas_chart, totales_chart = [], []
    
    hoy = timezone.now().date()
    if dias == 365:
        ventas_agrupadas = ventas_validas.annotate(mes=TruncMonth('fecha')).values('mes').annotate(total_mes=Sum('total')).order_by('mes')
        datos_dict = {v['mes'].strftime('%Y-%m'): float(v['total_mes']) for v in ventas_agrupadas if v['mes']}
        
        meses_es = {1: 'Ene', 2: 'Feb', 3: 'Mar', 4: 'Abr', 5: 'May', 6: 'Jun', 7: 'Jul', 8: 'Ago', 9: 'Sep', 10: 'Oct', 11: 'Nov', 12: 'Dic'}
        
        for i in range(11, -1, -1):
            m = hoy.month - i
            y = hoy.year
            while m <= 0:
                m += 12
                y -= 1
                
            llave = f"{y}-{m:02d}"
            fechas_chart.append(f"{meses_es[m]} {y}")
            totales_chart.append(datos_dict.get(llave, 0.0))
    else:
        ventas_diarias = ventas_validas.values('fecha__date').annotate(total_dia=Sum('total')).order_by('fecha__date')
        fechas_dict = {v['fecha__date']: float(v['total_dia']) for v in ventas_diarias if v['fecha__date']}
        
        for i in range(dias - 1, -1, -1):
            fecha_iter = hoy - timedelta(days=i)
            fechas_chart.append(fecha_iter.strftime('%d %b'))
            totales_chart.append(fechas_dict.get(fecha_iter, 0.0))
        
    # 4. Top 10 Productos Más Exitosos
    top_productos = DetalleVenta.objects.filter(venta__empresa=empresa, venta__fecha__gte=fecha_inicio).exclude(venta__estado='anulada')\
        .values('producto__nombre', 'producto__tipo_producto__nombre')\
        .annotate(cantidad_vendida=Count('id'), ingresos=Sum('precio_unitario')).order_by('-cantidad_vendida')[:10]
        
    # 5. Top 10 Mejores Clientes
    top_clientes = ventas_validas.filter(cliente__isnull=False)\
        .values('cliente__nombre', 'cliente__ruc_cedula')\
        .annotate(num_compras=Count('id'), total_comprado=Sum('total')).order_by('-total_comprado')[:10]

    top_rentabilidad = metricas_rentabilidad.get('top_rentabilidad', [])
    top_rotacion = metricas_rentabilidad.get('top_rotacion', [])
    categorias_rentables = metricas_rentabilidad.get('categorias_rentables', [])

    context = {
        'empresa': empresa, 'dias': dias, 'total_ingresos': total_ingresos, 'total_ventas': total_ventas,
        'ticket_promedio': ticket_promedio, 'top_productos': top_productos, 'top_clientes': top_clientes,
        'fechas_chart': json.dumps(fechas_chart), 'totales_chart': json.dumps(totales_chart),
        'es_dueno': es_dueno,
        'metricas_rentabilidad': metricas_rentabilidad,
        'rentabilidad_promedio': metricas_rentabilidad.get('resumen', {}).get('rentabilidad_promedio', 0),
        'utilidad_total': metricas_rentabilidad.get('resumen', {}).get('utilidad_total', 0),
        'producto_mayor_rentabilidad': metricas_rentabilidad.get('producto_mayor_rentabilidad'),
        'alta_rotacion_count': metricas_rentabilidad.get('alta_rotacion_count', 0),
        'top_productos_rentables': top_rentabilidad,
        'top_productos_rotacion': top_rotacion,
        'categorias_mas_rentables': categorias_rentables,
    }
    return render(request, 'dashboard/tenant_reportes.html', context)


def tienda_publica(request, slug):
    """Landing page / Catálogo público dinámico del cliente final"""
    # Buscamos la empresa sin importar si está activa o no
    empresa = get_object_or_404(Empresa, slug=slug)
    
    # Verificamos si la tienda está suspendida (por administración o por falta de pago)
    suspendida = False
    if not empresa.activa:
        suspendida = True
    elif hasattr(empresa, 'suscripcion') and empresa.suscripcion.dias_restantes() <= 0:
        suspendida = True
        
    if suspendida:
        return render(request, 'dashboard/tienda_publica.html', {'empresa': empresa, 'suspendida': True})

    # OPTIMIZACIÓN: Obtener todos los productos y ordenar para que los disponibles salgan primero
    productos = Producto.objects.filter(empresa=empresa).select_related('tipo_producto').order_by('estado', '-creado_en')
    
    # OPTIMIZACIÓN: Delegar la agrupación a la Base de Datos (Ahorro de CPU y RAM)
    categorias = Producto.objects.filter(empresa=empresa).exclude(categoria='').values_list('categoria', flat=True).distinct().order_by('categoria')
    
    context = {'empresa': empresa, 'productos': productos, 'categorias': categorias, 'suspendida': False}
    return render(request, 'dashboard/tienda_publica.html', context)


@login_required
def saas_equipo(request):
    """Gestión de administradores globales del sistema (SaaS)"""
    if getattr(request.user.profile, 'rol', '') != ROLE_SUPERADMIN:
        raise PermissionDenied("Acceso exclusivo para administradores globales.")
        
    # Evaluación Inteligente: Cierra votaciones expiradas (48h) o ejecuta si hay consenso antes
    for req_obj in AdminActionRequest.objects.filter(status='pending'):
        req_obj.check_and_execute()
        
    admins = UserProfile.objects.filter(rol=ROLE_SUPERADMIN).select_related('user').order_by('-creado_en')
    active_requests = AdminActionRequest.objects.filter(status='pending').select_related('target_admin__user', 'requester__user').prefetch_related('votes_approve', 'votes_reject')
    
    if request.method == 'POST':
        action = request.POST.get('action', 'crear')
        
        if action == 'crear':
            username = request.POST.get('username')
            email = request.POST.get('email')
            password = request.POST.get('password')
            
            if username and email and password:
                try:
                    with transaction.atomic():
                        # 1. Crear el usuario nativo de Django
                        user = User.objects.create_user(username=username, email=email, password=password)
                        user.is_staff = True  # Para que puedan acceder al /admin/ nativo si se requiere
                        user.save()
                        
                        # 2. Obtener o crear su perfil de forma segura
                        profile, _ = UserProfile.objects.get_or_create(user=user)
                        profile.rol = ROLE_SUPERADMIN
                        profile.onboarding_completado = True
                        profile.save()
                        
                        messages.success(request, f"Administrador {username} añadido exitosamente al Equipo del Sistema.")
                except Exception as e:
                    messages.error(request, f"No se pudo crear el administrador: Ya existe un usuario con ese nombre o correo.")
                    
        elif action in ['solicitar_suspension', 'solicitar_eliminacion']:
            admin_id = request.POST.get('admin_id')
            target_admin = get_object_or_404(UserProfile, id=admin_id, rol=ROLE_SUPERADMIN)
            
            if target_admin.user == request.user:
                messages.error(request, "Por políticas de seguridad, no puedes realizar esta acción sobre tu propia cuenta.")
                return redirect('saas_equipo')
                
            # Evaluar el consenso: Buscamos a TODOS los otros admins, excluyéndote a ti y al afectado
            otros_admins = UserProfile.objects.filter(rol=ROLE_SUPERADMIN).exclude(id__in=[target_admin.id, request.user.profile.id])
            
            if not otros_admins.exists():
                messages.warning(request, "Acción denegada: Se requiere al menos un tercer administrador en el sistema para aprobar esta solicitud (Protocolo de Consenso de Dueños).")
            else:
                # Verificar que no exista un consenso pendiente para este usuario
                if AdminActionRequest.objects.filter(target_admin=target_admin, status='pending').exists():
                    messages.warning(request, "Ya existe una votación en curso para este administrador.")
                    return redirect('saas_equipo')
                    
                accion_txt = "suspender" if action == 'solicitar_suspension' else "eliminar"
                
                # Crear la propuesta y agregar el voto "A favor" del creador automáticamente
                nueva_solicitud = AdminActionRequest.objects.create(target_admin=target_admin, requester=request.user.profile, action_type=action)
                nueva_solicitud.votes_approve.add(request.user.profile)
                nueva_solicitud.check_and_execute() # Por si solo existen 2 admins en total
                
                messages.success(request, f"Protocolo de Consenso activado: Se ha notificado en el panel al resto de dueños para autorizar '{accion_txt}' a '{target_admin.user.username}'. Tienen 48 horas para votar.")
                
        elif action == 'votar_consenso':
            req_id = request.POST.get('request_id')
            voto = request.POST.get('voto')
            req = get_object_or_404(AdminActionRequest, id=req_id, status='pending')
            
            if req.target_admin == request.user.profile:
                messages.error(request, "Seguridad: No puedes votar en tu propio protocolo de consenso.")
            elif req.votes_approve.filter(id=request.user.profile.id).exists() or req.votes_reject.filter(id=request.user.profile.id).exists():
                messages.error(request, "Ya has emitido tu voto para esta solicitud.")
            else:
                if voto == 'approve':
                    req.votes_approve.add(request.user.profile)
                else:
                    req.votes_reject.add(request.user.profile)
                messages.success(request, "Tu voto ha sido registrado correctamente.")
                req.check_and_execute() # Evaluamos si con este voto se cumple la meta
                
        return redirect('saas_equipo')
        
    return render(request, 'dashboard/saas_equipo.html', {'admins': admins, 'active_requests': active_requests})