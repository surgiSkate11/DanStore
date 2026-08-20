import openpyxl
import csv
import io
import json
import logging
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import IntegrityError
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.core.paginator import Paginator
from openpyxl.styles import Font, PatternFill, Alignment

from .models import Producto, TipoProducto
from .forms import ProductoForm
from empresas.mixins import requiere_modulo
from tiendita_project.utils import form_errors_to_messages

logger = logging.getLogger(__name__)

# ── UTILIDAD ──────────────────────────────────────────────────────────────────

def get_negocio_tipo(user):
    """Obtiene el tipo de negocio del perfil del usuario de forma segura"""
    try:
        return user.profile.negocio_tipo or 'general'
    except Exception:
        return 'general'


# ── VISTAS PRINCIPALES ────────────────────────────────────────────────────────
# Lista productos 

@login_required
@requiere_modulo('productos', accion='ver')
def lista_productos(request):
    """Lista productos con filtros genéricos para cualquier negocio"""

    # select_related evita N+1 queries
    # CRÍTICO: Filtrar siempre por la empresa del request actual
    productos = Producto.objects.filter(empresa=request.empresa).select_related('tipo_producto').only(
        'id', 'imagen', 'nombre', 'marca', 'color', 'categoria',
        'variante', 'precio', 'estado', 'descripcion',
        'variante', 'precio', 'precio_costo', 'estado', 'descripcion',
        'tipo_producto__nombre', 'stock'
    )

    # ── FILTROS ───────────────────────────────────────────────────────────────
    # Categoría libre
    categoria = request.GET.get('categoria', '').strip()
    if categoria:
        productos = productos.filter(categoria__icontains=categoria)

    # Tipo de producto
    tipo_id = request.GET.get('tipo_producto', '').strip()
    if tipo_id and tipo_id.isdigit():
        productos = productos.filter(tipo_producto_id=tipo_id)

    # Estado
    estado = request.GET.get('estado', '').strip()
    if estado and estado in dict(Producto.ESTADO_CHOICES):
        productos = productos.filter(estado=estado)

    # Búsqueda general — busca en nombre, descripción, marca, categoría, variante
    busqueda = request.GET.get('busqueda', '').strip()
    if busqueda:
        productos = productos.filter(
            Q(nombre__icontains=busqueda) |
            Q(descripcion__icontains=busqueda) |
            Q(marca__icontains=busqueda) |
            Q(categoria__icontains=busqueda) |
            Q(variante__icontains=busqueda) |
            Q(color__icontains=busqueda)
        )

    # Paginación: 52 por página (divisible en 4 columnas)
    paginator = Paginator(productos, 52)
    page_obj = paginator.get_page(request.GET.get('page'))

    # Para el filtro de categorías — muestra las que existen en la BD
    categorias_existentes = (
        Producto.objects.filter(empresa=request.empresa).exclude(categoria='')
        .values_list('categoria', flat=True)
        .distinct()
        .order_by('categoria')
    )

    tipos_existentes = TipoProducto.objects.filter(empresa=request.empresa)

    context = {
        'page_obj': page_obj,
        'categorias': categorias_existentes,
        'tipos': tipos_existentes,
        'estados': Producto.ESTADO_CHOICES,
        # Valores actuales de filtros para mantener selección
        'filtro_categoria': categoria,
        'filtro_tipo': tipo_id,
        'filtro_estado': estado,
        'filtro_busqueda': busqueda,
    }
    return render(request, 'productos/lista_productos.html', context)


@login_required
@requiere_modulo('productos', accion='crear')
def crear_producto(request):
    """Crea un nuevo producto — genérico para cualquier negocio"""
    if request.method == 'POST':
        form = ProductoForm(request.POST, request.FILES, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, '¡Producto creado exitosamente!')
            return redirect('lista_productos')
        else:
            form_errors_to_messages(request, form)
    else:
        form = ProductoForm(empresa=request.empresa)

    return render(request, 'productos/form_producto.html', {
        'form': form,
        'titulo': 'Crear Nuevo Producto',
    })


@login_required
@requiere_modulo('productos', accion='editar')
def editar_producto(request, pk):
    """Edita un producto existente"""
    producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)

    if request.method == 'POST':
        form = ProductoForm(request.POST, request.FILES, instance=producto, empresa=request.empresa)
        if form.is_valid():
            form.save()
            messages.success(request, '¡Producto actualizado exitosamente!')
            return redirect('lista_productos')
        else:
            form_errors_to_messages(request, form)
    else:
        form = ProductoForm(instance=producto, empresa=request.empresa)

    return render(request, 'productos/form_producto.html', {
        'form': form,
        'titulo': 'Editar Producto',
        'producto': producto,
    })


@login_required
@requiere_modulo('productos', accion='eliminar')
def eliminar_producto(request, pk):
    """Elimina un producto con confirmación"""
    producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)

    if request.method == 'POST':
        producto.delete()
        messages.success(request, '¡Producto eliminado exitosamente!')
        return redirect('lista_productos')

    return render(request, 'productos/confirmar_eliminar.html', {'producto': producto})


# ── VISTAS AJAX ───────────────────────────────────────────────────────────────

@login_required
@requiere_modulo('productos', accion='editar')
def cambiar_estado_producto(request, pk):
    """AJAX — cambia el estado de un producto"""
    if request.method == 'POST':
        producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)
        nuevo_estado = request.POST.get('estado')

        if nuevo_estado in dict(Producto.ESTADO_CHOICES):
            producto.estado = nuevo_estado
            producto.save(update_fields=['estado', 'actualizado_en'])
            return JsonResponse({
                'success': True,
                'estado': producto.get_estado_display()
            })

    return JsonResponse({'success': False}, status=400)


@login_required
@requiere_modulo('productos', accion='editar')
def cambiar_categoria_producto(request, pk):
    """AJAX — cambia la categoría de un producto (texto libre)"""
    if request.method == 'POST':
        producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)
        nueva_categoria = request.POST.get('categoria', '').strip()
        producto.categoria = nueva_categoria
        producto.save(update_fields=['categoria', 'actualizado_en'])
        return JsonResponse({'success': True, 'categoria': producto.categoria})

    return JsonResponse({'success': False}, status=400)


@login_required
@requiere_modulo('productos', accion='crear')
def api_crear_tipo_producto(request):
    """AJAX — crea un nuevo Tipo de Producto desde un modal"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            nombre = data.get('nombre', '').strip()
        except:
            nombre = request.POST.get('nombre', '').strip()
            
        if not nombre:
            return JsonResponse({'success': False, 'error': 'El nombre no puede estar vacío.'}, status=400)
            
        # Comprobar si ya existe
        if TipoProducto.objects.filter(nombre__iexact=nombre, empresa=request.empresa).exists():
            return JsonResponse({'success': False, 'error': 'Ya existe un tipo de producto con ese nombre.'}, status=400)
            
        # Crear
        nuevo_tipo = TipoProducto.objects.create(nombre=nombre, empresa=request.empresa)
        return JsonResponse({
            'success': True,
            'id': nuevo_tipo.id,
            'nombre': nuevo_tipo.nombre
        })
    return JsonResponse({'success': False, 'error': 'Método no permitido'}, status=405)


@login_required
@requiere_modulo('productos', accion='editar')
def cambiar_variante_producto(request, pk):
    """
    AJAX — cambia la variante de un producto.
    Reemplaza cambiar_talla_producto — ahora es genérico:
    sirve para talla, modelo, sabor, versión, etc.
    """
    if request.method == 'POST':
        producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)
        nueva_variante = request.POST.get('variante', '').strip()
        producto.variante = nueva_variante
        producto.save(update_fields=['variante', 'actualizado_en'])
        return JsonResponse({'success': True, 'variante': producto.variante})

    return JsonResponse({'success': False}, status=400)


@login_required
@requiere_modulo('productos', accion='editar')
def subir_foto_producto(request, pk):
    """AJAX — sube o reemplaza la foto de un producto"""
    if request.method == 'POST' and request.FILES.get('imagen'):
        try:
            producto = get_object_or_404(Producto, pk=pk, empresa=request.empresa)
            producto.imagen = request.FILES['imagen']
            producto.save(update_fields=['imagen', 'actualizado_en'])
            return JsonResponse({
                'success': True,
                'message': 'Foto actualizada exitosamente',
                'imagen_url': producto.imagen.url if producto.imagen else None
            })
        except Exception as e:
            logger.exception('Error al subir foto del producto %s (empresa %s)', pk, getattr(request.empresa, 'id', None))
            return JsonResponse({'success': False, 'message': 'No se pudo guardar la foto. Intenta nuevamente.'}, status=400)

    return JsonResponse({'success': False, 'message': 'No se recibió imagen'}, status=400)


@login_required
@requiere_modulo('productos', accion='ver')
def get_categorias_sugeridas(request):
    """
    AJAX — devuelve las categorías existentes para autocompletar.
    Reemplaza get_rangos_edad que era específico de ropa.
    """
    q = request.GET.get('q', '').strip()
    categorias = (
        Producto.objects.filter(empresa=request.empresa).exclude(categoria='')
        .filter(categoria__icontains=q)
        .values_list('categoria', flat=True)
        .distinct()
        .order_by('categoria')[:20]
    )
    return JsonResponse({'categorias': list(categorias)})


# ── IMPORTACIÓN EXCEL / CSV ───────────────────────────────────────────────────

@login_required
@requiere_modulo('productos', accion='crear')
def importar_excel(request):
    """
    Importa productos desde Excel (.xlsx) o CSV.

    Columnas del archivo (fila 1 = cabecera, se ignora):
      A: Nombre del producto     (texto, OBLIGATORIO)
      B: Categoría               (texto libre, opcional)
      C: Precio                  (número, OBLIGATORIO)
      D: Estado                  (disponible / separado / vendido)
      E: Descripción             (texto, opcional)
      F: Variante                (talla / modelo / sabor / etc., opcional)
      G: Marca                   (texto, opcional)
      H: Stock                   (número entero, opcional, default 1)
      I: Precio de Oferta        (número, opcional — vacío = sin oferta)
      J: Color                   (texto, opcional)

    Si el nombre ya existe → actualiza. Si es nuevo → crea.
    """
    if request.method == 'POST':
        archivo = request.FILES.get('archivo')

        if not archivo:
            messages.error(request, 'Debes seleccionar un archivo.')
            return redirect('importar_excel')

        # VALIDACIÓN: usuario sin empresa asignada. Sin este guard, cualquier
        # fila intenta guardarse con empresa=None y Postgres revienta con un
        # IntegrityError crudo (null value in column "empresa_id"...) que
        # antes se mostraba tal cual al usuario.
        if not getattr(request, 'empresa', None):
            logger.warning('Intento de importar Excel sin empresa asignada. Usuario: %s', request.user.id)
            messages.error(request, 'No se encontró una empresa asociada a tu cuenta. Contacta al administrador.')
            return redirect('importar_excel')

        if not request.empresa.activa:
            messages.error(request, 'Tu cuenta está suspendida. No puedes importar productos en este momento.')
            return redirect('importar_excel')

        if archivo.size == 0:
            messages.error(request, 'El archivo está vacío.')
            return redirect('importar_excel')

        nombre_archivo = archivo.name.lower()
        errores = []
        creados = 0
        actualizados = 0

        try:
            # ── LEER ARCHIVO ──────────────────────────────────────────────────
            if nombre_archivo.endswith('.xlsx') or nombre_archivo.endswith('.xls'):
                try:
                    wb = openpyxl.load_workbook(archivo)
                except Exception:
                    logger.exception('Error al abrir el archivo Excel. Usuario: %s', request.user.id)
                    messages.error(request, 'El archivo no es un Excel válido o está dañado. Verifica el formato.')
                    return redirect('importar_excel')
                ws = wb.active
                filas = list(ws.iter_rows(min_row=2, values_only=True))

            elif nombre_archivo.endswith('.csv'):
                try:
                    contenido = archivo.read().decode('utf-8-sig')
                except UnicodeDecodeError:
                    logger.warning('CSV con codificación inválida. Usuario: %s', request.user.id)
                    messages.error(request, 'El archivo CSV tiene una codificación no compatible. Guárdalo como UTF-8 e intenta de nuevo.')
                    return redirect('importar_excel')
                reader = csv.reader(io.StringIO(contenido))
                try:
                    next(reader)  # saltar cabecera
                except StopIteration:
                    messages.error(request, 'El archivo está vacío.')
                    return redirect('importar_excel')
                filas = list(reader)

            else:
                messages.error(request, 'Formato no soportado. Usa .xlsx o .csv')
                return redirect('importar_excel')

            if not filas:
                messages.warning(request, 'El archivo no contiene filas de datos para importar.')
                return redirect('importar_excel')

            # ── PROCESAR FILAS ────────────────────────────────────────────────
            for i, fila in enumerate(filas, start=2):
                try:
                    # Extraer valores con defaults seguros
                    nombre_prod   = str(fila[0]).strip() if len(fila) > 0 and fila[0] else None
                    categoria     = str(fila[1]).strip() if len(fila) > 1 and fila[1] else ''
                    precio        = _parsear_decimal(fila[2]) if len(fila) > 2 else None
                    estado        = str(fila[3]).strip().lower() if len(fila) > 3 and fila[3] else 'disponible'
                    descripcion   = str(fila[4]).strip() if len(fila) > 4 and fila[4] else ''
                    variante      = str(fila[5]).strip() if len(fila) > 5 and fila[5] else ''
                    marca         = str(fila[6]).strip() if len(fila) > 6 and fila[6] else ''
                    stock         = int(float(str(fila[7]))) if len(fila) > 7 and fila[7] else 1
                    precio_oferta = _parsear_decimal(fila[8]) if len(fila) > 8 else None
                    color         = str(fila[9]).strip() if len(fila) > 9 and fila[9] else ''

                    # Fila completamente vacía — saltar silenciosamente
                    if not nombre_prod:
                        continue

                    # 'precio' ahora es obligatorio a nivel de base de datos
                    # (ver migración 0005). Si la fila no trae un precio
                    # válido, se reporta como error legible en vez de dejar
                    # que reviente con un IntegrityError feo al guardar.
                    if precio is None:
                        errores.append(f'Fila {i}: falta el precio (es obligatorio) para "{nombre_prod}"')
                        continue

                    # Mapeo inteligente (Opción Flexible SaaS)
                    mapeo_estados = {
                        'disponible': 'disponible', 'stock': 'disponible', 'libre': 'disponible',
                        'reservado': 'reservado', 'separado': 'reservado', 'apartado': 'reservado', 'pendiente': 'reservado',
                        'vendido': 'vendido', 'vendida': 'vendido', 'pagado': 'vendido', 'entregado': 'vendido'
                    }

                    # Normalizar estado buscando coincidencias, default a 'disponible'
                    estado_final = 'disponible'
                    for clave, valor_oficial in mapeo_estados.items():
                        if clave in estado:
                            estado_final = valor_oficial
                            break

                    # Precio de oferta: solo se guarda si es realmente menor al
                    # precio normal (mismo criterio que valida el formulario
                    # manual de productos) — evita "ofertas" más caras que el
                    # precio normal coladas por error de tipeo en el Excel.
                    if precio_oferta is not None and precio is not None and precio_oferta >= precio:
                        precio_oferta = None

                    # update_or_create por nombre — evita duplicados
                    producto, creado = Producto.objects.update_or_create(
                        nombre=nombre_prod,
                        empresa=request.empresa,
                        defaults={
                            'categoria': categoria,
                            'precio': precio,
                            'precio_oferta': precio_oferta,
                            'estado': estado_final,
                            'descripcion': descripcion,
                            'variante': variante,
                            'marca': marca,
                            'color': color,
                            'stock': stock,
                            'origen': 'excel',
                        }
                    )

                    if creado:
                        creados += 1
                    else:
                        actualizados += 1

                except (IndexError, ValueError, TypeError) as e:
                    errores.append(f'Fila {i}: revisa que los datos tengan el formato correcto.')
                    logger.info('Error de formato en fila %s al importar Excel (empresa %s): %s', i, request.empresa.id, e)

                except IntegrityError as e:
                    # No exponer el mensaje crudo de la base de datos al usuario.
                    errores.append(f'Fila {i}: no se pudo guardar "{nombre_prod}" — faltan datos obligatorios o hay un conflicto con un registro existente.')
                    logger.error('IntegrityError en fila %s al importar Excel (empresa %s): %s', i, request.empresa.id, e)

        except Exception as e:
            # Cualquier error no anticipado: se registra completo para diagnóstico
            # interno, pero al usuario solo se le muestra un mensaje genérico.
            logger.exception('Error inesperado al importar Excel (empresa %s, usuario %s)', getattr(request.empresa, 'id', None), request.user.id)
            messages.error(request, 'Ocurrió un error inesperado al procesar el archivo. Intenta nuevamente o contacta al administrador.')
            return redirect('importar_excel')

        # ── RESULTADO ─────────────────────────────────────────────────────────
        if creados or actualizados:
            messages.success(request, f'✅ {creados} productos creados, {actualizados} actualizados.')
        if errores:
            for error in errores[:5]:
                messages.warning(request, error)
            if len(errores) > 5:
                messages.warning(request, f'... y {len(errores) - 5} errores más.')
        if not creados and not actualizados and not errores:
            messages.warning(request, 'No se encontraron filas válidas para importar.')

        return redirect('lista_productos')

    return render(request, 'productos/importar_excel.html')


def _parsear_decimal(valor):
    """Convierte un valor a decimal de forma segura — maneja comas y símbolos"""
    if valor is None:
        return None
    texto = str(valor).strip().replace(',', '.').replace('$', '').replace(' ', '')
    if not texto:
        return None
    try:
        return float(texto)
    except ValueError:
        return None


@login_required
@requiere_modulo('productos', accion='crear')
def descargar_plantilla(request):
    """
    Descarga una plantilla Excel con las columnas correctas y una hoja de referencia.
    Genérica — sirve para cualquier tipo de negocio.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    # ── CABECERAS ─────────────────────────────────────────────────────────────
    cabeceras = ['Nombre', 'Categoría', 'Precio', 'Estado', 'Descripción', 'Variante', 'Marca', 'Stock', 'Precio de Oferta', 'Color']
    ws.append(cabeceras)

    # ── FILAS DE EJEMPLO por tipo de negocio ──────────────────────────────────
    ejemplos = [
        # Ropa (con oferta de ejemplo)
        ['Vestido floral', 'Mujer', '25.00', 'disponible', 'Talla M, color rojo', 'M', 'Zara', 1, '19.99', 'Rojo'],
        # Tecnología
        ['iPhone 15 Pro', 'Smartphones', '999.00', 'disponible', '128GB, color negro', '128GB', 'Apple', 2, '', 'Negro'],
        # Alimentos
        ['Torta de chocolate', 'Postres', '15.00', 'disponible', 'Para 10 personas', '10 porciones', 'Casera', 1, '', ''],
        # Servicios
        ['Diseño de logo', 'Servicios', '80.00', 'reservado', 'Seña pagada por cliente', 'Paquete básico', '', 1, '', ''],
    ]
    for ejemplo in ejemplos:
        ws.append(ejemplo)

    # ── ESTILO CABECERAS ──────────────────────────────────────────────────────
    color_cabecera = PatternFill("solid", fgColor="1a73e8")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = color_cabecera
        cell.alignment = Alignment(horizontal='center')

    # ── ESTILO FILAS EJEMPLO ──────────────────────────────────────────────────
    color_ejemplo = PatternFill("solid", fgColor="EBF5FB")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.fill = color_ejemplo

    # ── ANCHO DE COLUMNAS ─────────────────────────────────────────────────────
    anchos = {'A': 30, 'B': 18, 'C': 12, 'D': 15, 'E': 35, 'F': 18, 'G': 15, 'H': 8, 'I': 16, 'J': 12}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    # ── HOJA DE REFERENCIA ────────────────────────────────────────────────────
    ws_ref = wb.create_sheet("Referencia")
    ws_ref.title = "Referencia"

    ws_ref.append(['Campo', 'Descripción', 'Valores válidos / Ejemplos'])
    ref_data = [
        ['Nombre',      'Nombre del producto (obligatorio)',           'Vestido, iPhone 15, Torta de chocolate'],
        ['Categoría',   'Categoría libre — escribe la tuya',          'Mujer, Electrónica, Postres, Servicios'],
        ['Precio',      'Precio de venta en números',                  '25.00 / 999 / 15.50'],
        ['Estado',      'Estado actual del producto (flexible)',       'disponible | reservado (o separado/apartado) | vendido'],
        ['Descripción', 'Descripción adicional (opcional)',            'Talla M, color rojo / 128GB negro'],
        ['Variante',    'Talla, modelo, sabor, versión (opcional)',    'M / L / 128GB / Chocolate / Básico'],
        ['Marca',       'Marca o proveedor (opcional)',                'Zara / Apple / Casera'],
        ['Stock',       'Cantidad disponible (opcional, default 1)',   '1 / 5 / 10'],
    ]
    for fila in ref_data:
        ws_ref.append(fila)

    # Estilo hoja referencia
    for cell in ws_ref[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="27AE60")
        cell.alignment = Alignment(horizontal='center')

    ws_ref.column_dimensions['A'].width = 15
    ws_ref.column_dimensions['B'].width = 38
    ws_ref.column_dimensions['C'].width = 42

    # ── RESPUESTA ─────────────────────────────────────────────────────────────
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="plantilla_productos.xlsx"'
    wb.save(response)
    return response